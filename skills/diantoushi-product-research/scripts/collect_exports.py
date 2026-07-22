#!/usr/bin/env python3
import argparse
import json
import re
import shutil
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Optional


EXPORT_PATTERNS = {
    "product_data": "商品数据ID_*.xlsx",
    "sku_preview": "店透-SKU预览-表格-*.xlsx",
    "ask_all": "店透视-问大家分析-*.xlsx",
}


def safe_name(value: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|\n\r\t]+', "-", value).strip(" .-")
    return cleaned[:60] or "商品"


def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    stem, suffix = path.stem, path.suffix
    for index in range(1, 1000):
        candidate = path.with_name(f"{stem} ({index}){suffix}")
        if not candidate.exists():
            return candidate
    raise RuntimeError(f"Could not find available filename for {path}")


def workbook_info(path: Path) -> dict:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        return {
            "valid_xlsx": True,
            "sheet": ws.title,
            "rows": ws.max_row,
            "columns": ws.max_column,
            "headers": headers,
        }
    except Exception as exc:
        return {"valid_xlsx": False, "error": str(exc)}


def find_exports(download_dir: Path, item_id: Optional[str], since: float) -> list[tuple[str, Path]]:
    found: list[tuple[str, Path]] = []
    for kind, pattern in EXPORT_PATTERNS.items():
        for path in download_dir.glob(pattern):
            if item_id and item_id not in path.name:
                continue
            if path.stat().st_mtime < since:
                continue
            found.append((kind, path))
    return sorted(found, key=lambda entry: entry[1].stat().st_mtime)


def main() -> int:
    parser = argparse.ArgumentParser(description="Collect 店透视 export workbooks into a Desktop folder.")
    parser.add_argument("--product-name", required=True)
    parser.add_argument("--item-id")
    parser.add_argument("--since-epoch", type=float, default=time.time() - 3600)
    parser.add_argument("--download-dir", default=str(Path.home() / "Downloads"))
    parser.add_argument("--desktop-dir", default=str(Path.home() / "Desktop"))
    parser.add_argument("--copy", action="store_true", help="Copy instead of moving files.")
    args = parser.parse_args()

    download_dir = Path(args.download_dir).expanduser()
    desktop_dir = Path(args.desktop_dir).expanduser()
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    folder_parts = ["店透视导出", safe_name(args.product_name)]
    if args.item_id:
        folder_parts.append(args.item_id)
    folder_parts.append(timestamp)
    target_dir = desktop_dir / "-".join(folder_parts)
    target_dir.mkdir(parents=True, exist_ok=True)

    exports = find_exports(download_dir, args.item_id, args.since_epoch)
    moved = []
    for kind, source in exports:
        destination = unique_path(target_dir / source.name)
        if args.copy:
            shutil.copy2(source, destination)
            action = "copied"
        else:
            shutil.move(str(source), str(destination))
            action = "moved"
        moved.append(
            {
                "kind": kind,
                "action": action,
                "source": str(source),
                "destination": str(destination),
                "workbook": workbook_info(destination),
            }
        )

    summary = {
        "target_dir": str(target_dir),
        "product_name": args.product_name,
        "item_id": args.item_id,
        "since_epoch": args.since_epoch,
        "files": moved,
        "missing_kinds": [kind for kind in EXPORT_PATTERNS if not any(file["kind"] == kind for file in moved)],
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if moved else 2


if __name__ == "__main__":
    sys.exit(main())
